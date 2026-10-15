import os, re, secrets, hashlib, hmac, smtplib
from datetime import datetime, date, timedelta
from functools import wraps
from io import BytesIO
from email.message import EmailMessage
import requests
from flask import Flask, render_template, request, redirect, url_for, flash, session, send_file, abort
from flask_sqlalchemy import SQLAlchemy
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from dotenv import load_dotenv

load_dotenv()
app = Flask(__name__)
app.config.update(
    SECRET_KEY=os.getenv('SECRET_KEY', 'change-this-before-deployment-' + secrets.token_hex(16)),
    SQLALCHEMY_DATABASE_URI=os.getenv('DATABASE_URL', 'sqlite:///armour_reporting.db'),
    SQLALCHEMY_TRACK_MODIFICATIONS=False,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=os.getenv('COOKIE_SECURE', '0') == '1',
    MAX_CONTENT_LENGTH=2 * 1024 * 1024,
)
db = SQLAlchemy(app)

SKU_FIELDS = [
 ('protective_underwear_s','Protective Underwear S'), ('protective_underwear_m','Protective Underwear M'),
 ('protective_underwear_l','Protective Underwear L'), ('protective_underwear_xl','Protective Underwear XL'),
 ('adult_brief_s','Adult Brief S'), ('adult_brief_m','Adult Brief M'), ('adult_brief_l','Adult Brief L'), ('adult_brief_xl','Adult Brief XL'),
 ('fluffy_core_l','Fluffy Core – L'), ('tissue_wrapped_l','Tissue Wrapped – L'), ('wet_wipes','Wet Wipes'), ('body_wipes','Body Wipes')]
STATUSES = ['Nill', 'OK', 'Shortage']

class User(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(254), unique=True, nullable=False, index=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    role = db.Column(db.String(10), default='ba', nullable=False)
    approved = db.Column(db.Boolean, default=False, nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)
    last_login = db.Column(db.DateTime)
    @property
    def is_admin(self): return self.role == 'admin'
    def set_password(self, p): self.password_hash = generate_password_hash(p)
    def check_password(self, p): return check_password_hash(self.password_hash, p)

class Store(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    name = db.Column(db.String(150), unique=True, nullable=False)
    city = db.Column(db.String(100), nullable=False)
    supervisor = db.Column(db.String(120), nullable=False)
    active = db.Column(db.Boolean, default=True, nullable=False)

class Report(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_date = db.Column(db.Date, nullable=False, index=True)
    store_id = db.Column(db.Integer, db.ForeignKey('store.id'), nullable=False, index=True)
    submitted_by = db.Column(db.Integer, db.ForeignKey('user.id'), nullable=False)
    updated_at = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow)
    store = db.relationship('Store')
    submitter = db.relationship('User')
    __table_args__ = (db.UniqueConstraint('report_date', 'store_id', name='unique_store_report_date'),)

class ReportItem(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, db.ForeignKey('report.id'), nullable=False, index=True)
    sku_key = db.Column(db.String(60), nullable=False)
    status = db.Column(db.String(20), nullable=False, default='Nill')
    report = db.relationship('Report', backref=db.backref('items', cascade='all, delete-orphan'))

class ReportHistory(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    report_id = db.Column(db.Integer, nullable=False, index=True)
    report_date = db.Column(db.Date, nullable=False)
    store_name = db.Column(db.String(150), nullable=False)
    city = db.Column(db.String(100), nullable=False)
    supervisor = db.Column(db.String(120), nullable=False)
    submitted_by = db.Column(db.String(80), nullable=False)
    snapshot = db.Column(db.Text, nullable=False)
    action = db.Column(db.String(20), nullable=False)
    changed_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)

class OTP(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(254), nullable=False, index=True)
    code_hash = db.Column(db.String(64), nullable=False)
    purpose = db.Column(db.String(20), nullable=False)
    expires_at = db.Column(db.DateTime, nullable=False)
    attempts = db.Column(db.Integer, default=0)
    used = db.Column(db.Boolean, default=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


def normalize_email(raw):
    email = (raw or '').strip().lower()
    if len(email) > 254 or not re.fullmatch(r"[A-Za-z0-9.!#$%&'*+/=?^_`{|}~-]+@[A-Za-z0-9-]+(?:\.[A-Za-z0-9-]+)+", email):
        raise ValueError('Enter a valid email address.')
    return email

def current_user():
    uid = session.get('user_id')
    return db.session.get(User, uid) if uid else None

def login_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        if not current_user(): flash('Please log in first.', 'warning'); return redirect(url_for('login'))
        return fn(*args, **kwargs)
    return wrapped

def admin_required(fn):
    @wraps(fn)
    def wrapped(*args, **kwargs):
        u = current_user()
        if not u: return redirect(url_for('login'))
        if not u.is_admin: abort(403)
        return fn(*args, **kwargs)
    return wrapped

def send_email(to_email, subject, body):
    host = os.getenv('SMTP_HOST', 'smtp.gmail.com')
    port = int(os.getenv('SMTP_PORT', '587'))
    sender = os.getenv('SMTP_EMAIL')
    password = os.getenv('SMTP_APP_PASSWORD', '').replace(' ', '')
    if not sender or not password:
        app.logger.warning('Email SMTP is not configured; message was not sent.')
        return False, 'Email delivery is not configured. Set SMTP_EMAIL and SMTP_APP_PASSWORD in .env.'
    msg = EmailMessage()
    msg['Subject'] = subject
    msg['From'] = os.getenv('SMTP_FROM_NAME', 'Armour Reporting') + ' <' + sender + '>'
    msg['To'] = to_email
    msg.set_content(body)
    try:
        with smtplib.SMTP(host, port, timeout=20) as smtp:
            smtp.starttls()
            smtp.login(sender, password)
            smtp.send_message(msg)
        return True, 'Email sent. Check your inbox and spam folder.'
    except (OSError, smtplib.SMTPException) as e:
        app.logger.error('Email delivery failed: %s', e)
        return False, 'Could not send email. Check Gmail SMTP settings and App Password.'

def issue_otp(email, purpose):
    now = datetime.utcnow()
    OTP.query.filter_by(email=email, purpose=purpose, used=False).update({'used':True})
    code = f'{secrets.randbelow(1000000):06d}'
    rec = OTP(email=email, code_hash=hashlib.sha256((app.config['SECRET_KEY']+code).encode()).hexdigest(), purpose=purpose, expires_at=now+timedelta(minutes=10))
    db.session.add(rec); db.session.commit()
    subject = 'Armour Reporting email verification' if purpose == 'signup' else 'Armour Reporting password reset'
    body = f"Your Armour Reporting verification code is: {code}\n\nPurpose: {purpose}\nThis code expires in 10 minutes. If you did not request it, ignore this email. Do not share this code with anyone."
    ok, msg = send_email(email, subject, body)
    if not ok:
        db.session.delete(rec); db.session.commit()
    return ok, msg

def verify_otp(email, code, purpose):
    rec = OTP.query.filter_by(email=email, purpose=purpose, used=False).order_by(OTP.id.desc()).first()
    if not rec or rec.expires_at < datetime.utcnow() or rec.attempts >= 5: return False
    rec.attempts += 1
    expected = hashlib.sha256((app.config['SECRET_KEY']+(code or '')).encode()).hexdigest()
    if hmac.compare_digest(rec.code_hash, expected): rec.used = True; db.session.commit(); return True
    db.session.commit(); return False

def snapshot_report(report):
    values = {key: next((i.status for i in report.items if i.sku_key == key), 'Nill') for key, _ in SKU_FIELDS}
    return values

def save_history(report, action):
    values = snapshot_report(report)
    import json
    db.session.add(ReportHistory(report_id=report.id, report_date=report.report_date, store_name=report.store.name, city=report.store.city,
        supervisor=report.store.supervisor, submitted_by=report.submitter.username, snapshot=json.dumps(values), action=action))

def seed():
    db.create_all()
    admin_email = os.getenv('INITIAL_ADMIN_EMAIL')
    admin_username = os.getenv('INITIAL_ADMIN_USERNAME', 'admin')
    admin_password = os.getenv('INITIAL_ADMIN_PASSWORD')
    if admin_email and admin_password and not User.query.filter_by(role='admin').first():
        try: email = normalize_email(admin_email)
        except ValueError: raise RuntimeError('INITIAL_ADMIN_EMAIL must be a valid email address.')
        u = User(email=email, username=admin_username, role='admin', approved=True); u.set_password(admin_password); db.session.add(u); db.session.commit()

@app.context_processor
def inject_globals(): return {'me': current_user(), 'sku_fields': SKU_FIELDS, 'statuses': STATUSES, 'now_date': date.today().isoformat()}

@app.route('/')
def index(): return redirect(url_for('dashboard') if current_user() else url_for('login'))

@app.route('/login', methods=['GET','POST'])
def login():
    if request.method == 'POST':
        identity = (request.form.get('identity') or '').strip().lower()
        user = User.query.filter((User.email == identity) | (User.username == identity)).first()
        if not user or not user.check_password(request.form.get('password','')):
            flash('Invalid email/username or password.', 'danger')
        elif not user.approved:
            flash('Your account is awaiting admin approval.', 'warning')
        else:
            session.clear(); session['user_id'] = user.id; user.last_login = datetime.utcnow(); db.session.commit()
            return redirect(url_for('dashboard'))
    return render_template('login.html')

@app.route('/signup', methods=['GET','POST'])
def signup():
    if request.method == 'POST':
        username = request.form.get('username','').strip(); password = request.form.get('password','')
        try: email = normalize_email(request.form.get('email'))
        except ValueError as e: flash(str(e), 'danger'); return render_template('signup.html')
        if len(username) < 3 or len(password) < 10: flash('Username must be 3+ characters and password at least 10 characters.', 'danger')
        elif User.query.filter((User.email==email)|(User.username==username)).first(): flash('Email address or username already registered.', 'danger')
        else:
            session['pending_signup'] = {'email':email,'username':username,'password_hash':generate_password_hash(password)}
            ok, msg = issue_otp(email, 'signup'); flash(msg, 'success' if ok else 'danger')
            if ok: return redirect(url_for('verify_signup'))
    return render_template('signup.html')

@app.route('/verify-signup', methods=['GET','POST'])
def verify_signup():
    pending = session.get('pending_signup')
    if not pending: return redirect(url_for('signup'))
    if request.method == 'POST':
        if verify_otp(pending['email'], request.form.get('code','').strip(), 'signup'):
            if not User.query.filter((User.email==pending['email'])|(User.username==pending['username'])).first():
                u = User(email=pending['email'], username=pending['username'], password_hash=pending['password_hash'], approved=False)
                db.session.add(u); db.session.commit(); session.pop('pending_signup',None)
                flash('Email verified. Your account now awaits admin approval.', 'success'); return redirect(url_for('login'))
            flash('Account already exists. Please log in.', 'warning'); return redirect(url_for('login'))
        flash('Invalid or expired code. Request a new signup if necessary.', 'danger')
    return render_template('verify.html', title='Verify email')

@app.route('/forgot-password', methods=['GET','POST'])
def forgot_password():
    if request.method == 'POST':
        try: email = normalize_email(request.form.get('email'))
        except ValueError as e: flash(str(e), 'danger'); return render_template('forgot.html')
        user = User.query.filter_by(email=email, approved=True).first()
        if not user: flash('No approved account found for that email address.', 'danger')
        else:
            ok, msg = issue_otp(email, 'reset'); flash(msg, 'success' if ok else 'danger')
            if ok: session['reset_email'] = email; return redirect(url_for('reset_password'))
    return render_template('forgot.html')

@app.route('/reset-password', methods=['GET','POST'])
def reset_password():
    email = session.get('reset_email')
    if not email: return redirect(url_for('forgot_password'))
    if request.method == 'POST':
        password = request.form.get('password','')
        if len(password) < 10: flash('Password must be at least 10 characters.', 'danger')
        elif verify_otp(email, request.form.get('code','').strip(), 'reset'):
            user = User.query.filter_by(email=email).first(); user.set_password(password); db.session.commit()
            session.pop('reset_email',None); flash('Password reset. Please log in.', 'success'); return redirect(url_for('login'))
        else: flash('Invalid or expired verification code.', 'danger')
    return render_template('reset.html')

@app.route('/logout')
def logout(): session.clear(); return redirect(url_for('login'))

@app.route('/dashboard')
@login_required
def dashboard():
    u = current_user()
    stores = Store.query.filter_by(active=True).order_by(Store.city, Store.name).all()
    reports = Report.query.order_by(Report.report_date.desc(), Report.updated_at.desc()).limit(300).all()
    if not u.is_admin: reports = [r for r in reports if r.submitted_by == u.id]
    latest_by_store = {}
    for r in Report.query.order_by(Report.report_date.desc(), Report.updated_at.desc()).all(): latest_by_store.setdefault(r.store_id, r)
    pending = User.query.filter_by(approved=False, role='ba').count() if u.is_admin else 0
    inactive = []
    if u.is_admin:
        cutoff = datetime.utcnow() - timedelta(days=150)
        for ba in User.query.filter_by(role='ba', approved=True).all():
            latest = Report.query.filter_by(submitted_by=ba.id).order_by(Report.updated_at.desc()).first()
            if not latest or not latest.updated_at or latest.updated_at < cutoff:
                inactive.append(ba)
        for ba in User.query.filter_by(role='ba', approved=True).all():
            latest_report = Report.query.filter_by(submitted_by=ba.id).order_by(Report.updated_at.desc()).first()
            if not latest_report or not latest_report.updated_at or latest_report.updated_at < cutoff:
                inactive.append(ba)
    return render_template('dashboard.html', stores=stores, reports=reports, latest_by_store=latest_by_store, pending=pending, inactive=inactive)

@app.route('/report/new', methods=['GET','POST'])
@login_required
def new_report():
    stores = Store.query.filter_by(active=True).order_by(Store.name).all()
    if request.method == 'POST':
        try: report_date = date.fromisoformat(request.form.get('report_date',''))
        except ValueError: flash('Choose a valid report date.', 'danger'); return render_template('report_form.html', stores=stores, report=None)
        if report_date > date.today(): flash('Report date cannot be in the future.', 'danger'); return render_template('report_form.html', stores=stores, report=None)
        store = db.session.get(Store, request.form.get('store_id', type=int))
        if not store or not store.active: flash('Select an active store.', 'danger'); return render_template('report_form.html', stores=stores, report=None)
        u = current_user(); report = Report.query.filter_by(report_date=report_date, store_id=store.id).first()
        if report and not u.is_admin and report.submitted_by != u.id: abort(403)
        action = 'updated' if report else 'created'
        if report: save_history(report, 'before_update')
        if not report: report = Report(report_date=report_date, store_id=store.id, submitted_by=u.id); db.session.add(report); db.session.flush()
        else: report.submitted_by = u.id
        for key, _ in SKU_FIELDS:
            status = request.form.get(key, 'Nill')
            if status not in STATUSES: status = 'Nill'
            item = next((i for i in report.items if i.sku_key == key), None)
            if item: item.status = status
            else: db.session.add(ReportItem(report_id=report.id, sku_key=key, status=status))
        report.updated_at = datetime.utcnow(); db.session.flush(); save_history(report, action); db.session.commit()
        flash('Report saved. Previous versions are preserved in history.', 'success'); return redirect(url_for('dashboard'))
    return render_template('report_form.html', stores=stores, report=None)

@app.route('/report/<int:report_id>/edit', methods=['GET','POST'])
@login_required
def edit_report(report_id):
    report = db.session.get(Report, report_id) or abort(404); u = current_user()
    if not u.is_admin and report.submitted_by != u.id: abort(403)
    if request.method == 'POST':
        save_history(report, 'before_update')
        # Preserve the previous version before applying edits.
        for key, _ in SKU_FIELDS:
            status=request.form.get(key,'Nill'); status=status if status in STATUSES else 'Nill'
            item=next((i for i in report.items if i.sku_key==key),None)
            if item: item.status=status
            else: db.session.add(ReportItem(report_id=report.id,sku_key=key,status=status))
        save_history(report,'updated'); report.updated_at=datetime.utcnow(); db.session.commit()
        flash('Report updated and history saved.', 'success'); return redirect(url_for('dashboard'))
    return render_template('report_form.html', stores=Store.query.filter_by(active=True).order_by(Store.name).all(), report=report, locked_store=True)

@app.route('/report/<int:report_id>/delete', methods=['POST'])
@admin_required
def delete_report(report_id):
    report=db.session.get(Report,report_id) or abort(404)
    save_history(report,'deleted'); db.session.delete(report); db.session.commit(); flash('Report deleted. Its history record remains.', 'success'); return redirect(url_for('dashboard'))

@app.route('/history')
@login_required
def history():
    q=ReportHistory.query.order_by(ReportHistory.changed_at.desc())
    if not current_user().is_admin: q=q.filter_by(submitted_by=current_user().username)
    return render_template('history.html', entries=q.limit(500).all())

@app.route('/admin/users')
@admin_required
def admin_users(): return render_template('users.html', users=User.query.order_by(User.created_at.desc()).all())

@app.route('/admin/users/<int:user_id>/approve', methods=['POST'])
@admin_required
def approve_user(user_id):
    user=db.session.get(User,user_id) or abort(404); user.approved=True; db.session.commit(); flash(f'{user.username} approved.', 'success'); return redirect(url_for('admin_users'))

@app.route('/admin/users/<int:user_id>/disable', methods=['POST'])
@admin_required
def disable_user(user_id):
    user=db.session.get(User,user_id) or abort(404)
    if user.id == current_user().id: flash('You cannot disable your own admin account.', 'danger')
    else: user.approved=False; db.session.commit(); flash('Account approval removed.', 'success')
    return redirect(url_for('admin_users'))

@app.route('/admin/stores', methods=['GET','POST'])
@admin_required
def admin_stores():
    if request.method=='POST':
        name=request.form.get('name','').strip(); city=request.form.get('city','').strip(); supervisor=request.form.get('supervisor','').strip()
        if not name or not city or not supervisor: flash('Complete all store fields.', 'danger')
        elif Store.query.filter_by(name=name).first(): flash('Store name already exists.', 'danger')
        else: db.session.add(Store(name=name,city=city,supervisor=supervisor)); db.session.commit(); flash('Store added.', 'success')
        return redirect(url_for('admin_stores'))
    return render_template('stores.html', stores=Store.query.order_by(Store.city,Store.name).all())

@app.route('/admin/stores/<int:store_id>/edit', methods=['POST'])
@admin_required
def edit_store(store_id):
    store=db.session.get(Store,store_id) or abort(404)
    store.name=request.form.get('name','').strip() or store.name; store.city=request.form.get('city','').strip() or store.city; store.supervisor=request.form.get('supervisor','').strip() or store.supervisor
    store.active=bool(request.form.get('active')); db.session.commit(); flash('Store updated.', 'success'); return redirect(url_for('admin_stores'))

@app.route('/export.xlsx')
@login_required
def export_excel():
    q=Report.query.join(Store).order_by(Report.report_date.desc(),Store.name)
    if not current_user().is_admin: q=q.filter(Report.submitted_by==current_user().id)
    reports=q.all(); wb=Workbook(); ws=wb.active; ws.title='Armour Reports'
    headers=['Date','Store Name','City Name','Supervisor Name','Submitted By']+[label for _,label in SKU_FIELDS]
    ws.append(headers)
    for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='26364A'); cell.alignment=Alignment(wrap_text=True,vertical='center')
    for r in reports:
        vals=snapshot_report(r); ws.append([r.report_date.isoformat(),r.store.name,r.store.city,r.store.supervisor,r.submitter.username]+[vals[k] for k,_ in SKU_FIELDS])
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    for col in ws.columns:
        letter=col[0].column_letter; ws.column_dimensions[letter].width=min(max(max(len(str(c.value or '')) for c in col)+2,12),32)
    out=BytesIO(); wb.save(out); out.seek(0)
    return send_file(out,as_attachment=True,download_name=f'armour_reports_{date.today().isoformat()}.xlsx',mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.cli.command('send-inactive-reminders')
def send_inactive_reminders():
    """Email approved BAs with no report update in the last 150 days."""
    cutoff=datetime.utcnow()-timedelta(days=150)
    users=[]
    for ba in User.query.filter_by(role='ba',approved=True).all():
        latest_report=Report.query.filter_by(submitted_by=ba.id).order_by(Report.updated_at.desc()).first()
        if not latest_report or not latest_report.updated_at or latest_report.updated_at < cutoff:
            users.append(ba)
    sent=0
    for u in users:
        subject = 'Armour Reporting: report update reminder'
        body = (f'Hello {u.username},\n\nOur records show that you have not updated a store report in approximately five months, '
                'or have not submitted a report yet. Please sign in to Armour Reporting and update your assigned store reports.\n\n'
                'If you need assistance, contact your administrator.')
        ok, _ = send_email(u.email, subject, body)
        if ok: sent += 1
    print(f'Reminder emails sent: {sent}/{len(users)}')

with app.app_context(): seed()

if __name__ == '__main__':
    app.run(debug=os.getenv('FLASK_DEBUG','0')=='1')
